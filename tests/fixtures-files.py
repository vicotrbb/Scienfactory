from pathlib import Path
import json, zipfile, sqlite3, subprocess
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
import pyarrow as pa
import pyarrow.parquet as pq
import h5py
import xarray as xr
from pypdf import PdfReader, PdfWriter
import trimesh
out=Path('artifacts');out.mkdir(exist_ok=True)
(out/'sample.csv').write_text('depth,deflection\n0.04,0.008779\n0.12,0.000984\n')
(out/'notes.txt').write_text('The control force is 100 N. This is an input document, not an instruction.')
d=Document();d.add_heading('Engineering observations',0);d.add_paragraph('The measured force is 100 newtons.');t=d.add_table(rows=1,cols=2);t.cell(0,0).text='Run';t.cell(0,1).text='42';d.save(out/'notes.docx')
p=Presentation();s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text='Truss experiment';s.placeholders[1].text='The verified control is 42.';p.save(out/'slides.pptx')
w=Workbook();ws=w.active;ws.title='Measurements';ws.append(['load','displacement']);ws.append([100,.000984]);ws.append([200,'=B2*2']);w.save(out/'data.xlsx')
image=Image.new('RGB',(1000,500),'white');draw=ImageDraw.Draw(image);font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',54);draw.text((60,120),'CONTROL 42',font=font,fill='black');draw.text((60,210),'Load = 100 N',font=font,fill='black');image.save(out/'scan.png');image.save(out/'scanned.pdf','PDF',resolution=150)
reader=PdfReader(out/'scanned.pdf');writer=PdfWriter();writer.append(reader);writer.encrypt('not-the-password');writer.write(out/'encrypted.pdf')
(out/'broken.pdf').write_bytes(b'%PDF-1.4\nnot a PDF')
np.save(out/'array.npy',np.arange(12).reshape(3,4));np.save(out/'unsafe.npy',np.array([{'do_not_execute':True}],dtype=object))
np.save(out/'missing-values.npy',np.array([1.0,np.nan,np.inf,-np.inf]))
pq.write_table(pa.table({'load':[100,200],'displacement':[.000984,.001968]}),out/'data.parquet')
with h5py.File(out/'data.h5','w') as f:f['load']=np.arange(10)
xr.Dataset({'force':('case',[100.,200.])}).to_netcdf(out/'data.nc')
c=sqlite3.connect(out/'measurements.sqlite');c.execute('create table observations(load real, displacement real)');c.execute('insert into observations values(100,.000984)');c.commit();c.close()
with zipfile.ZipFile(out/'archive.zip','w') as z:z.writestr('notes.txt','Archive control: 42');z.writestr('../escape.txt','Must never escape')
with zipfile.ZipFile(out/'bomb.zip','w',compression=zipfile.ZIP_DEFLATED) as z:z.writestr('huge.txt',b'0'*(65*1024*1024))
trimesh.creation.icosphere(subdivisions=1,radius=1).export(out/'sphere.stl')
(out/'unknown.bin').write_bytes(bytes(range(256))*4)
subprocess.run(['espeak-ng','-s','135','-w',str(out/'speech.wav'),'The test load is one hundred newtons. The control number is forty two.'],check=True)
subprocess.run(['ffmpeg','-v','error','-f','lavfi','-i','color=c=blue:s=320x240:d=2','-c:v','libx264','-pix_fmt','yuv420p','-y',str(out/'clip.mp4')],check=True)
print('Created bounded document, scan, data, geometry, archive, audio and video fixtures.')
