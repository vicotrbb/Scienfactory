"""Bounded, offline file inspection. Parsers run only inside the disposable lab."""
import base64
import csv
import io
import json
import math
import pathlib
import sqlite3
import subprocess
import tarfile
import zipfile

MAX_TEXT = 24000
MAX_ARCHIVE = 64 * 1024 * 1024
OUT = pathlib.Path('artifacts')


def command(args, timeout=25):
    result = subprocess.run(args, capture_output=True, timeout=timeout)
    if result.returncode:
        raise ValueError(result.stderr.decode('utf8', 'replace')[-1500:])
    return result.stdout


def archive_guard(path):
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 3000 or sum(e.file_size for e in entries) > MAX_ARCHIVE:
                raise ValueError('Archive expansion exceeds the 64 MB / 3000 member inspection limit.')
            if any(e.flag_bits & 1 for e in entries):
                raise ValueError('Encrypted archive requires a decrypted copy.')
            return entries
    return None


def inspect(path, start=0, count=4, transcribe=True):
    from PIL import Image, UnidentifiedImageError
    Image.MAX_IMAGE_PIXELS = 30_000_000
    path = pathlib.Path(path)
    OUT.mkdir(exist_ok=True)
    count = min(max(int(count), 1), 4)
    start = max(int(start), 0)
    result = {'name': path.name, 'bytes': path.stat().st_size, 'status': 'inspected',
              'text': '', 'metadata': {}, 'previews': [], 'warnings': [], 'coverage': {}}
    def preview(p, mime, pages=1):
        result['previews'].append({'name': p.name, 'mime': mime, 'pages': pages})
    def image_preview(image, name):
        image = image.convert('RGB')
        image.thumbnail((1600, 1600))
        target = OUT / name
        image.save(target, 'PNG', optimize=True)
        preview(target, 'image/png')
        return target
    try:
        magic = command(['file', '-b', '--mime-type', str(path)], 5).decode().strip()
        result['detectedMime'] = magic
        ext = path.suffix.lower()
        entries = archive_guard(path)
        if magic == 'application/pdf' or ext == '.pdf':
            from pypdf import PdfReader, PdfWriter
            reader = PdfReader(path)
            if reader.is_encrypted and not reader.decrypt(''):
                raise ValueError('Encrypted PDF requires a decrypted copy.')
            total = len(reader.pages)
            selected = range(start, min(start + count, total))
            if start >= total:
                raise ValueError(f'Page offset exceeds {total} pages.')
            writer = PdfWriter()
            texts = []
            for index in selected:
                page = reader.pages[index]
                writer.add_page(page)
                text = page.extract_text() or ''
                target = OUT / f'page-{index+1:03d}'
                command(['pdftoppm', '-f', str(index+1), '-l', str(index+1), '-singlefile',
                         '-scale-to', '1600', '-png', str(path), str(target)])
                image_path = target.with_suffix('.png')
                preview(image_path, 'image/png')
                if len(text.strip()) < 40:
                    text = command(['tesseract', str(image_path), 'stdout'], 15).decode('utf8', 'replace')
                    result['warnings'].append(f'Page {index+1}: OCR text can contain recognition errors.')
                texts.append(f'Page {index+1}\n{text}')
            target = OUT / 'selected-pages.pdf'
            with target.open('wb') as stream:
                writer.write(stream)
            preview(target, 'application/pdf', len(selected))
            result['text'] = '\n\n'.join(texts)
            result['coverage'] = {'unit': 'pages', 'start': start, 'count': len(selected), 'total': total}
            result['metadata'] = {'pages': total, 'document': {str(k): str(v) for k,v in (reader.metadata or {}).items()}}
        elif magic.startswith('image/') and ext != '.svg':
            with Image.open(path) as image:
                total = getattr(image, 'n_frames', 1)
                result['metadata'] = {'width': image.width, 'height': image.height, 'mode': image.mode, 'frames': total}
                for index in range(start, min(start+count, total)):
                    image.seek(index)
                    target = image_preview(image.copy(), f'image-{index+1:03d}.png')
                    result['text'] += command(['tesseract', str(target), 'stdout'], 15).decode('utf8', 'replace')
                result['coverage'] = {'unit': 'frames', 'start': start, 'count': min(count, max(total-start, 0)), 'total': total}
                result['warnings'].append('OCR is a fallible text aid. Inspect the visual attachment for diagrams and layout.')
        elif ext in ('.docx', '.docm'):
            from docx import Document
            doc = Document(path)
            blocks = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                blocks.extend(' | '.join(c.text for c in row.cells) for row in table.rows[:100])
            result['text'] = '\n'.join(blocks)[start:start+MAX_TEXT]
            result['coverage'] = {'unit': 'characters', 'start': start, 'total': sum(len(b)+1 for b in blocks)}
            result['warnings'].append('Extracted paragraphs and tables; layout, embedded objects, and tracked-change interpretation are not certified.')
        elif ext == '.doc':
            result['text'] = command(['antiword', str(path)]).decode('utf8', 'replace')[start:start+MAX_TEXT]
        elif ext in ('.pptx', '.pptm'):
            from pptx import Presentation
            deck = Presentation(path)
            blocks = []
            for index in range(start, min(start+count, len(deck.slides))):
                slide = deck.slides[index]
                lines = [shape.text for shape in slide.shapes if shape.has_text_frame]
                for shape in slide.shapes:
                    if shape.has_table:
                        lines += [' | '.join(c.text for c in row.cells) for row in shape.table.rows]
                    if hasattr(shape, 'image') and len(result['previews']) < 4:
                        image_preview(Image.open(io.BytesIO(shape.image.blob)), f'slide-{index+1}-image-{len(result["previews"])+1}.png')
                if slide.has_notes_slide:
                    lines.append(slide.notes_slide.notes_text_frame.text)
                blocks.append(f'Slide {index+1}\n' + '\n'.join(lines))
            result['text'] = '\n\n'.join(blocks)
            result['coverage'] = {'unit': 'slides', 'start': start, 'count': min(count, max(len(deck.slides)-start, 0)), 'total': len(deck.slides)}
            result['warnings'].append('Extracted slide text, tables, notes, and selected embedded images; original slide layout is not reconstructed.')
        elif ext in ('.xlsx', '.xlsm', '.xls'):
            if ext == '.xls':
                import xlrd
                book = xlrd.open_workbook(path)
                sheets = [{'name': s.name, 'rows': s.nrows, 'columns': s.ncols,
                           'sample': [s.row_values(i)[:30] for i in range(start, min(start+100, s.nrows))]} for s in book.sheets()[:20]]
            else:
                from openpyxl import load_workbook
                book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
                sheets = [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column,
                           'sample': [list(row) for row in s.iter_rows(min_row=start+1, max_row=min(start+100, s.max_row or 100), max_col=min(30, s.max_column or 30), values_only=True)]} for s in book.worksheets[:20]]
                book.close()
            result['metadata'] = {'sheets': sheets}
            result['text'] = json.dumps(sheets, default=str)
            result['coverage'] = {'unit': 'rows per sheet', 'start': start, 'count': 100, 'maxColumns': 30, 'maxSheets': 20}
            result['warnings'].append('Formulas are read as stored; they are not recalculated or executed.')
        elif ext in ('.odt', '.ods', '.odp', '.epub'):
            from bs4 import BeautifulSoup
            with zipfile.ZipFile(path) as archive:
                names = [n for n in archive.namelist() if n == 'content.xml' or n.endswith(('.xhtml', '.html'))]
                result['text'] = '\n'.join(BeautifulSoup(archive.read(n), 'html.parser').get_text(' ', strip=True) for n in names[:50])[start:start+MAX_TEXT]
            result['warnings'].append('Text extraction; original layout and interactive content are not rendered.')
        elif ext in ('.parquet', '.feather', '.arrow'):
            import pyarrow.parquet as pq
            import pyarrow.feather as feather
            table = pq.read_table(path) if ext == '.parquet' else feather.read_table(path)
            result['metadata'] = {'rows': table.num_rows, 'schema': str(table.schema)}
            result['text'] = json.dumps(table.slice(start, 50).to_pylist(), default=str)
            result['coverage'] = {'unit': 'rows', 'start': start, 'count': 50, 'total': table.num_rows}
        elif ext in ('.npy', '.npz'):
            import numpy as np
            data = np.load(path, allow_pickle=False)
            arrays = {n: data[n] for n in data.files[:20]} if ext == '.npz' else {'array': data}
            result['metadata'] = {n: {'shape': a.shape, 'dtype': str(a.dtype), 'sample': a.reshape(-1)[start:start+100].tolist()} for n,a in arrays.items()}
            result['text'] = json.dumps(result['metadata'], default=str)
        elif ext in ('.h5', '.hdf5', '.hdf'):
            import h5py
            with h5py.File(path, 'r') as file:
                datasets = []
                def visit(name, obj):
                    if len(datasets) >= 100:
                        return 'limit'
                    if isinstance(obj, h5py.Dataset):
                        datasets.append({'name': name, 'shape': obj.shape, 'dtype': str(obj.dtype), 'virtual': obj.is_virtual})
                file.visititems(visit)
            result['metadata'] = {'datasets': datasets}
            result['warnings'].append('Dataset structure inspected; use numerical code for explicitly selected bounded slices. External data is not loaded.')
        elif ext in ('.nc', '.netcdf'):
            import xarray as xr
            with xr.open_dataset(path, decode_times=False) as ds:
                result['metadata'] = {'dimensions': dict(ds.sizes), 'variables': {k: {'dimensions': v.dims, 'shape': v.shape, 'dtype': str(v.dtype)} for k,v in ds.variables.items()}}
        elif ext in ('.sqlite', '.sqlite3', '.db') and path.read_bytes()[:16] == b'SQLite format 3\x00':
            connection = sqlite3.connect(f'{path.resolve().as_uri()}?mode=ro&immutable=1', uri=True)
            connection.set_progress_handler(lambda: 1, 100000)
            tables = connection.execute("SELECT name,sql FROM sqlite_master WHERE type='table' LIMIT 50").fetchall()
            result['metadata'] = {'tables': tables}
            result['text'] = json.dumps(tables)
            connection.close()
        elif ext in ('.glb', '.gltf', '.stl', '.obj', '.ply', '.off', '.3mf'):
            import trimesh
            model = trimesh.load(path, force='scene')
            meshes = [g for g in model.geometry.values() if isinstance(g, trimesh.Trimesh)]
            result['metadata'] = {'meshes': len(meshes), 'bounds': model.bounds.tolist(), 'vertices': sum(len(m.vertices) for m in meshes), 'faces': sum(len(m.faces) for m in meshes)}
            import matplotlib
            matplotlib.use('Agg')
            import matplotlib.pyplot as plt
            fig = plt.figure(figsize=(7, 5), layout='constrained')
            ax = fig.add_subplot(projection='3d')
            for mesh in meshes[:10]:
                faces = mesh.faces[::max(1, len(mesh.faces)//3000)]
                ax.plot_trisurf(mesh.vertices[:,0], mesh.vertices[:,1], mesh.vertices[:,2], triangles=faces, color='#6b8c70', alpha=.9, linewidth=.1)
            ax.set(xlabel='x', ylabel='y', zlabel='z', title=path.name)
            target = OUT/'geometry-preview.png'
            fig.savefig(target, dpi=140)
            plt.close(fig)
            preview(target, 'image/png')
            result['warnings'].append('Geometry units and coordinate conventions must be confirmed from the source.')
        elif magic.startswith(('audio/', 'video/')) or ext in ('.mp3', '.wav', '.flac', '.m4a', '.ogg', '.mp4', '.mov', '.webm', '.mkv'):
            metadata = json.loads(command(['ffprobe', '-v', 'error', '-show_format', '-show_streams', '-of', 'json', str(path)]))
            result['metadata'] = metadata
            duration = float(metadata.get('format', {}).get('duration', 0))
            streams = metadata.get('streams', [])
            if any(s.get('codec_type') == 'video' for s in streams):
                for index in range(count):
                    at = start + index * max(1, min(45, max(duration-start,1))/count)
                    if duration and at >= duration:
                        break
                    target = OUT/f'frame-{at:.1f}s.png'
                    command(['ffmpeg', '-v', 'error', '-ss', str(at), '-i', str(path), '-frames:v', '1', '-vf', 'scale=1280:720:force_original_aspect_ratio=decrease', '-y', str(target)])
                    if target.exists():
                        preview(target, 'image/png')
            if any(s.get('codec_type') == 'audio' for s in streams) and transcribe:
                target = pathlib.Path('/tmp/inspection-audio.wav')
                command(['ffmpeg', '-v', 'error', '-ss', str(start), '-i', str(path), '-t', '45', '-ar', '16000', '-ac', '1', '-y', str(target)])
                from faster_whisper import WhisperModel
                model = WhisperModel('/opt/lab/whisper-base', device='cpu', compute_type='int8', cpu_threads=2, local_files_only=True)
                segments, info = model.transcribe(str(target), beam_size=1, condition_on_previous_text=False)
                result['text'] = '\n'.join(f'[{s.start+start:.1f}-{s.end+start:.1f}s] {s.text.strip()}' for s in segments)
                result['metadata']['transcriptLanguage'] = info.language
                result['warnings'].append('Automatic Whisper-base transcript is fallible, especially for technical terms, silence, or noisy recordings. Verify against the recording.')
            result['coverage'] = {'unit': 'seconds', 'start': start, 'count': min(45, max(duration-start, 0)), 'total': duration}
        elif entries is not None or tarfile.is_tarfile(path):
            if entries is not None:
                members = [{'name': e.filename, 'bytes': e.file_size, 'unsafePath': pathlib.PurePosixPath(e.filename).is_absolute() or '..' in pathlib.PurePosixPath(e.filename).parts} for e in entries[:500]]
            else:
                with tarfile.open(path) as archive:
                    members = []
                    for e in archive:
                        if len(members) >= 500:
                            break
                        members.append({'name': e.name, 'bytes': e.size, 'link': e.issym() or e.islnk(), 'unsafePath': pathlib.PurePosixPath(e.name).is_absolute() or '..' in pathlib.PurePosixPath(e.name).parts})
            result['metadata'] = {'members': members}
            result['text'] = json.dumps(members)
            result['warnings'].append('Archive inventory only. No member was extracted or executed. Select a safe member with inspect_file member to inspect its contents.')
        else:
            raw = path.read_bytes()
            encoding = 'utf-16' if raw[:2] in (b'\xff\xfe', b'\xfe\xff') else 'utf-8-sig'
            text = raw.decode(encoding, 'replace')
            if '\x00' not in text[:10000] and text[:10000].count('\ufffd') < max(1, len(text[:10000]) * .02):
                result['text'] = text[start:start+MAX_TEXT]
                result['coverage'] = {'unit': 'characters', 'start': start, 'count': len(result['text']), 'total': len(text)}
                if ext in ('.html', '.htm', '.xml', '.svg'):
                    from bs4 import BeautifulSoup
                    result['text'] = BeautifulSoup(text, 'html.parser').get_text(' ', strip=True)[start:start+MAX_TEXT]
                    result['warnings'].append('Markup inspected as inert text; scripts were not executed.')
            else:
                result['status'] = 'unsupported'
                result['metadata'] = {'headerHex': raw[:128].hex(), 'detectedMime': magic}
                result['warnings'].append('No installed decoder recognized this binary format. Original bytes are preserved. Supply a format specification or export to an open format; do not infer its contents from the filename.')
    except UnidentifiedImageError as error:
        result['status'] = 'unreadable' if path.suffix.lower() in ('.png', '.jpg', '.jpeg', '.tif', '.tiff', '.gif', '.webp', '.bmp') else 'unsupported'
        result['warnings'].append('The initial format guess was not confirmed by an image decoder. ' + str(error)[:1000])
    except Exception as error:
        result['status'] = 'unreadable'
        result['warnings'].append(str(error)[:1500])
    if len(result['text']) > MAX_TEXT:
        result['text'] = result['text'][:MAX_TEXT]
        result['warnings'].append('Extracted text truncated at 24,000 characters. Inspect a later range for more.')
    def finite_json(value):
        if isinstance(value, float) and not math.isfinite(value):
            return 'NaN' if math.isnan(value) else ('Infinity' if value > 0 else '-Infinity')
        if isinstance(value, dict):
            return {k: finite_json(v) for k, v in value.items()}
        if isinstance(value, (list, tuple)):
            return [finite_json(v) for v in value]
        return value
    result = finite_json(result)
    (OUT/'inspection.json').write_text(json.dumps(result, default=str, allow_nan=False))
    return result


def extract_member(path, member):
    """Return a bounded regular archive member, never write an archive-provided path."""
    target = pathlib.Path('/tmp') / ('member' + pathlib.Path(member).suffix[:16])
    pure = pathlib.PurePosixPath(member)
    if pure.is_absolute() or '..' in pure.parts or '\\' in member:
        raise ValueError('Unsafe archive member path.')
    if zipfile.is_zipfile(path):
        archive_guard(path)
        with zipfile.ZipFile(path) as archive:
            entry = archive.getinfo(member)
            if entry.is_dir() or entry.file_size > 8*1024*1024 or (entry.external_attr >> 16) & 0o170000 == 0o120000:
                raise ValueError('Only regular archive members up to 8 MB are supported.')
            data = archive.read(entry)
    else:
        with tarfile.open(path) as archive:
            entry = archive.getmember(member)
            if not entry.isfile() or entry.size > 8*1024*1024:
                raise ValueError('Only regular archive members up to 8 MB are supported.')
            data = archive.extractfile(entry).read(8*1024*1024+1)
    target.write_bytes(data)
    return target
